import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_cell_text(cell, text, bold=False, italic=False, font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT, color_rgb=(0x22, 0x22, 0x22)):
    cell.text = text
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if len(p.runs) > 0:
        run = p.runs[0]
        run.bold = bold
        run.italic = italic
        run.font.name = 'Arial'
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(*color_rgb)

def update_table_impact():
    doc_path = r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe.docx'
    doc = docx.Document(doc_path)
    
    # Table 7 es la tabla de evaluación de impacto de incidentes
    t3 = doc.tables[7]
    
    # Limpiamos filas existentes excepto el encabezado
    while len(t3.rows) > 1:
        tr = t3.rows[-1]._tr
        tr.getparent().remove(tr)
        
    # Re-poblamos exactamente con la estructura de la plantilla:
    # Header row 0
    format_cell_text(t3.rows[0].cells[0], "CASO", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xFF, 0xFF, 0xFF))
    format_cell_text(t3.rows[0].cells[1], "DESCRIPCIÓN", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xFF, 0xFF, 0xFF))
    format_cell_text(t3.rows[0].cells[2], "RESULTADO", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xFF, 0xFF, 0xFF))
    format_cell_text(t3.rows[0].cells[3], "IMPACTO", bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xFF, 0xFF, 0xFF))
    for c in t3.rows[0].cells:
        set_cell_background(c, "003366")
        set_cell_margins(c, 100, 100, 120, 120)

    # 1. CATEGORÍA: PRUEBAS FUNCIONALES
    r_cat1 = t3.add_row()
    r_cat1.cells[0].merge(r_cat1.cells[3])
    format_cell_text(r_cat1.cells[0], "PRUEBAS FUNCIONALES (VALIDACIONES Y CONTROL DE FLUJO)", bold=True, font_size=9, color_rgb=(0x00, 0x33, 0x66))
    set_cell_background(r_cat1.cells[0], "EBF2F8")
    set_cell_margins(r_cat1.cells[0], 80, 80, 100, 100)

    # Fila CP-02
    r_f1 = t3.add_row()
    format_cell_text(r_f1.cells[0], "CP-02", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_f1.cells[1], "Ingreso de paciente con RUT sin dígito verificador legítimo (11111111-9). Validación Módulo 11 en servidor.", font_size=8.5)
    format_cell_text(r_f1.cells[2], "Mitigado / Bloqueado. API rechaza con HTTP 400 'RUT inválido' y no contamina la base de datos.", font_size=8.5)
    format_cell_text(r_f1.cells[3], "LEVE", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0x2E, 0x8B, 0x57))
    for c in r_f1.cells: set_cell_margins(c)

    # Fila CP-03
    r_f2 = t3.add_row()
    format_cell_text(r_f2.cells[0], "CP-03", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_f2.cells[1], "Intento de eliminación de examen médico sin confirmación explícita. El ERS exige baja lógica protegida.", font_size=8.5)
    format_cell_text(r_f2.cells[2], "Mitigado / Controlado. API exige payload { confirm: true }, impidiendo borrado accidental de fichas.", font_size=8.5)
    format_cell_text(r_f2.cells[3], "MODERADO", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xB8, 0x86, 0x0B))
    for c in r_f2.cells: set_cell_margins(c)

    # 2. CATEGORÍA: PRUEBAS NO FUNCIONALES (SEGURIDAD Y CALIDAD - SONARQUBE)
    r_cat2 = t3.add_row()
    r_cat2.cells[0].merge(r_cat2.cells[3])
    format_cell_text(r_cat2.cells[0], "PRUEBAS NO FUNCIONALES (SEGURIDAD Y CALIDAD - SONARQUBE)", bold=True, font_size=9, color_rgb=(0x00, 0x33, 0x66))
    set_cell_background(r_cat2.cells[0], "EBF2F8")
    set_cell_margins(r_cat2.cells[0], 80, 80, 100, 100)

    # Fila SEG-01 (Hotspot CORS)
    r_nf1 = t3.add_row()
    format_cell_text(r_nf1.cells[0], "SEG-01\n(S5122)", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_nf1.cells[1], "Configuración de CORS permisivo con origin: '*' en app.js:21. En producción debe acotarse al dominio web del centro de salud.", font_size=8.5)
    format_cell_text(r_nf1.cells[2], "Revisado. Permitido en entorno local de pruebas. En producción se restringe a variable CORS_ORIGIN.", font_size=8.5)
    format_cell_text(r_nf1.cells[3], "MODERADO", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xB8, 0x86, 0x0B))
    for c in r_nf1.cells: set_cell_margins(c)

    # Fila SEG-02 (Hotspot Upload Limit)
    r_nf2 = t3.add_row()
    format_cell_text(r_nf2.cells[0], "SEG-02\n(S5693)", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_nf2.cells[1], "Control de tamaño máximo de subida (Multer 15MB) en upload.js:12. Riesgo de denegación de servicio (DoS) por memoria.", font_size=8.5)
    format_cell_text(r_nf2.cells[2], "Revisado / Mitigado. Límite acotado a 15MB por archivo con descarte automático por stream.", font_size=8.5)
    format_cell_text(r_nf2.cells[3], "MODERADO", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xB8, 0x86, 0x0B))
    for c in r_nf2.cells: set_cell_margins(c)

    # Fila SEG-03 (Hotspot Criptografía)
    r_nf3 = t3.add_row()
    format_cell_text(r_nf3.cells[0], "SEG-03\n(S2245)", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_nf3.cells[1], "Uso de generador pseudorandom Math.random() en sufijos de archivos clínicos en upload.js:18.", font_size=8.5)
    format_cell_text(r_nf3.cells[2], "Plan de mejora. Reemplazar por generador criptográfico fuerte crypto.randomBytes(16).toString('hex').", font_size=8.5)
    format_cell_text(r_nf3.cells[3], "LEVE", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0x2E, 0x8B, 0x57))
    for c in r_nf3.cells: set_cell_margins(c)

    # Fila CAL-01 (Smell Complejidad)
    r_nf4 = t3.add_row()
    format_cell_text(r_nf4.cells[0], "CAL-01\n(S3776)", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(r_nf4.cells[1], "Complejidad cognitiva de 76 en función query() de database.js:255 (umbral: 15). Deuda técnica de mantenibilidad.", font_size=8.5)
    format_cell_text(r_nf4.cells[2], "Plan de mejora. Modularizar consultas SQL emuladas en controladores por tabla específica.", font_size=8.5)
    format_cell_text(r_nf4.cells[3], "MODERADO", bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color_rgb=(0xB8, 0x86, 0x0B))
    for c in r_nf4.cells: set_cell_margins(c)

    # Guardar en las 3 ubicaciones
    doc.save(r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe.docx')
    doc.save(r'C:\Users\4ngel\.gemini\antigravity\scratch\cuidarte-plus\informe\EP3_ISY1102_Estudiante_Formato_Informe.docx')
    doc.save(r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe_Completado.docx')
    print("Tabla 5.3 actualizada con la estructura exacta de la plantilla!")

if __name__ == '__main__':
    update_table_impact()
