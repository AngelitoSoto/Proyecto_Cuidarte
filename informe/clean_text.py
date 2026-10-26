import docx

doc = docx.Document(r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe.docx')

doc.paragraphs[83].text = (
    "El ciclo de aseguramiento de calidad contemplo la ejecucion de 25 pruebas automatizadas "
    "de integracion y seguridad mediante Jest y Supertest, complementadas con el analisis estatico "
    "de 1.746 lineas de codigo fuente ejecutado por SonarQube LTS Community v9.9.8. El 100% de los casos "
    "disenados fueron ejecutados con exito, obteniendo estado 'OK' (Passed) en el Quality Gate oficial de SonarQube."
)

doc.paragraphs[84].text = (
    "• Cantidad Total de Casos de Prueba Disenados en ERS: 23 casos (100% de cobertura sobre RF-1 a RF-5 y NFRs).\n"
    "• Cantidad de Pruebas Automatizadas Ejecutadas: 25 pruebas en 5 suites (auth, patient, exam, user, security).\n"
    "• Resultados Obtenidos: 25 pruebas aprobadas (100% exito), 0 fallidas, 0 bloqueadas.\n"
    "• Metricas SonarQube: Quality Gate: PASSED (OK) | Bugs: 0 | Vulnerabilidades: 0 | Cobertura de Codigo: 70.7% | Duplicaciones: 0.0%."
)

for p_i in [85, 86]:
    doc.paragraphs[p_i].text = ""

doc.paragraphs[88].text = (
    "A continuacion se detalla la ejecucion caso a caso en la siguiente tabla, registrando los resultados "
    "especificos y las evidencias tecnicas obtenidas tanto de la bateria de pruebas automatizadas como de la auditoria estatica SonarQube:"
)

for p_i in [89, 90, 91, 92, 93, 94]:
    doc.paragraphs[p_i].text = ""

doc.paragraphs[96].text = (
    "Se adopto la escala estandar institucional para medir y clasificar el nivel de impacto de cualquier defecto o incidente "
    "detectado (Tabla de Escala de Criticidad). Para cada hallazgo identificado por SonarQube LTS Community (3 Security Hotspots y "
    "5 Code Smells) se determino su nivel de impacto y plan de remediacion tecnica:"
)

for p_i in [97, 98, 99]:
    doc.paragraphs[p_i].text = ""

# Guardar en todas las rutas
doc.save(r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe.docx')
doc.save(r'C:\Users\4ngel\.gemini\antigravity\scratch\cuidarte-plus\informe\EP3_ISY1102_Estudiante_Formato_Informe.docx')
doc.save(r'C:\Users\4ngel\Downloads\EP3_ISY1102_Estudiante_Formato_Informe_Completado.docx')

print("¡Informe limpiado y estructurado con éxito!")
